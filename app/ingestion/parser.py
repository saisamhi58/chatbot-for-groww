"""Document text extraction for multiple file types."""

import os
from typing import Optional
from dataclasses import dataclass, field


@dataclass
class ParsedDocument:
    """Result of parsing a document."""
    text: str
    pages: list[str] = field(default_factory=list)
    metadata: dict = field(default_factory=dict)


class DocumentParser:
    """Parse various document formats into plain text."""

    SUPPORTED_TYPES = {"pdf", "docx", "txt", "md", "html", "url"}

    def parse(self, file_path: str, file_type: str) -> ParsedDocument:
        """Parse a file based on its type. Use file_type='url' with a URL as file_path."""
        if file_type not in self.SUPPORTED_TYPES:
            raise ValueError(f"Unsupported file type: {file_type}")

        parser_method = getattr(self, f"_parse_{file_type}")
        return parser_method(file_path)

    def _parse_url(self, url: str) -> ParsedDocument:
        """Fetch a public web page and extract its text."""
        import requests
        from bs4 import BeautifulSoup

        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/120.0 Safari/537.36"
            )
        }
        response = requests.get(url, headers=headers, timeout=30)
        response.raise_for_status()

        soup = BeautifulSoup(response.text, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "header", "noscript"]):
            tag.decompose()

        title = soup.title.string.strip() if soup.title and soup.title.string else url
        text = soup.get_text(separator="\n", strip=True)

        return ParsedDocument(
            text=text,
            pages=[text],
            metadata={"source_url": url, "title": title},
        )

    def _parse_pdf(self, file_path: str) -> ParsedDocument:
        import pdfplumber

        pages = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""
                pages.append(text)

        return ParsedDocument(
            text="\n\n".join(pages),
            pages=pages,
            metadata={"page_count": len(pages)},
        )

    def _parse_docx(self, file_path: str) -> ParsedDocument:
        from docx import Document as DocxDocument

        doc = DocxDocument(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        return ParsedDocument(
            text="\n\n".join(paragraphs),
            pages=[],
            metadata={"paragraph_count": len(paragraphs)},
        )

    def _parse_txt(self, file_path: str) -> ParsedDocument:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()

        return ParsedDocument(
            text=text,
            pages=[text],
            metadata={},
        )

    def _parse_md(self, file_path: str) -> ParsedDocument:
        with open(file_path, "r", encoding="utf-8") as f:
            text = f.read()

        return ParsedDocument(
            text=text,
            pages=[text],
            metadata={},
        )

    def _parse_html(self, file_path: str) -> ParsedDocument:
        from bs4 import BeautifulSoup

        with open(file_path, "r", encoding="utf-8") as f:
            soup = BeautifulSoup(f.read(), "html.parser")

        # Remove script and style elements
        for script in soup(["script", "style"]):
            script.decompose()

        text = soup.get_text(separator="\n", strip=True)

        return ParsedDocument(
            text=text,
            pages=[text],
            metadata={},
        )
